"""Generate small binary/text fixtures under tests/fixtures/docs."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import (
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
)


def main() -> None:
    fx = Path(__file__).resolve().parents[1] / "tests" / "fixtures" / "docs"
    fx.mkdir(parents=True, exist_ok=True)

    (fx / "sample.txt").write_text(
        "Payment outage notes.\n\nThe billing service failed on 2024-03-12 after a deploy.",
        encoding="utf-8",
    )

    (fx / "sample.md").write_text(
        "---\n"
        "title: Incident Report March\n"
        "author: Alex Chen\n"
        "date: 2024-03-15\n"
        "---\n"
        "# Incident Report March\n\n"
        "## Timeline\n\n"
        "Dr. Smith noted the outage at 09:00 a.m. The API gateway returned 502 errors.\n\n"
        "## Impact\n\n"
        "Payment processing was delayed for two hours.\n",
        encoding="utf-8",
    )

    (fx / "sample.html").write_text(
        "<!DOCTYPE html>\n<html><head>\n"
        "<title>Change Log</title>\n"
        '<meta name="author" content="Ops Team" />\n'
        '<meta name="date" content="2024-03-01" />\n'
        "</head><body>\n"
        "<h1>Change Log</h1>\n"
        "<p>Deployed version 2.4.1 to production.</p>\n"
        "<h2>Services</h2>\n"
        "<p>Updated the payments worker and notification bus.</p>\n"
        "</body></html>\n",
        encoding="utf-8",
    )

    (fx / "sample.jsonl").write_text(
        '{"title": "Rec A", "author": "Ann", "date": "2024-01-01",'
        ' "text": "Alpha record about payments."}\n'
        '{"title": "Rec B", "text": "Beta record about outages."}\n',
        encoding="utf-8",
    )

    (fx / "sample.csv").write_text(
        "id,service,status\n1,payments,down\n2,notify,degraded\n",
        encoding="utf-8",
    )

    doc = Document()
    doc.core_properties.title = "DOCX Status"
    doc.core_properties.author = "Doc Author"
    doc.add_heading("DOCX Status", level=1)
    doc.add_paragraph("The ledger reconciler stalled after midnight.")
    doc.add_heading("Details", level=2)
    doc.add_paragraph("Operators restarted the worker at 01:15.")
    buf = BytesIO()
    doc.save(buf)
    (fx / "sample.docx").write_bytes(buf.getvalue())

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    content = b"BT /F1 12 Tf 72 720 Td (PDF sample about payment outages in March.) Tj ET"
    stream = DecodedStreamObject()
    stream.set_data(content)
    resources = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): DictionaryObject(
                        {
                            NameObject("/Type"): NameObject("/Font"),
                            NameObject("/Subtype"): NameObject("/Type1"),
                            NameObject("/BaseFont"): NameObject("/Helvetica"),
                        }
                    )
                }
            )
        }
    )
    page[NameObject("/Contents")] = stream
    page[NameObject("/Resources")] = resources
    writer.add_metadata({"/Title": "PDF Sample", "/Author": "PDF Author"})
    pdf_buf = BytesIO()
    writer.write(pdf_buf)
    (fx / "sample.pdf").write_bytes(pdf_buf.getvalue())

    empty = PdfWriter()
    empty.add_blank_page(width=612, height=792)
    empty_buf = BytesIO()
    empty.write(empty_buf)
    (fx / "empty.pdf").write_bytes(empty_buf.getvalue())

    print("wrote", sorted(p.name for p in fx.iterdir()))


if __name__ == "__main__":
    main()
