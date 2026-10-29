"""DOCX parser using python-docx."""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.errors import ParseError
from app.ingestion.metadata import apply_common_metadata, coalesce_str
from app.ingestion.normalize import normalize_text
from app.ingestion.parsers.base import ParsedDocument, Span, title_from_filename

_HEADING_STYLES = {
    "Heading 1": 1,
    "Heading 2": 2,
    "Heading 3": 3,
    "Heading 4": 4,
    "Heading 5": 5,
    "Heading 6": 6,
}


class DocxParser:
    """Parse Word ``.docx`` documents into text with heading spans."""

    def parse_bytes(self, data: bytes, *, filename: str | None = None) -> ParsedDocument:
        if not data:
            raise ParseError("DOCX document is empty")
        try:
            document = Document(io.BytesIO(data))
        except PackageNotFoundError as exc:
            raise ParseError(f"Invalid DOCX package: {exc}") from exc
        except Exception as exc:  # noqa: BLE001
            raise ParseError(f"Failed to read DOCX: {exc}") from exc

        blocks: list[tuple[str, str | None]] = []
        stack: list[str] = []
        for para in document.paragraphs:
            text = normalize_text(para.text or "")
            if not text:
                continue
            style_name = para.style.name if para.style is not None else ""
            level = _HEADING_STYLES.get(style_name)
            path: str | None
            if level is not None:
                stack = stack[: level - 1]
                stack.append(text)
                path = " > ".join(stack)
            else:
                path = " > ".join(stack) if stack else None
            blocks.append((text, path))

        if not blocks:
            raise ParseError("DOCX document has no extractable text")

        parts: list[str] = []
        spans: list[Span] = []
        offset = 0
        for block_text, path in blocks:
            if parts:
                parts.append("\n\n")
                offset += 2
            start = offset
            parts.append(block_text)
            offset += len(block_text)
            spans.append(Span(start=start, end=offset, heading_path=path))

        text = "".join(parts)
        core = _core_props(document)
        title = coalesce_str(core.get("title"), title_from_filename(filename)) or "untitled"
        author = coalesce_str(core.get("author"))
        metadata = apply_common_metadata(
            {},
            title=title,
            author=author,
            doc_date_raw=core.get("created"),
            source_type="docx",
        )
        return ParsedDocument(text=text, title=title, metadata=metadata, spans=spans)


def _core_props(document: Any) -> dict[str, Any]:
    props = document.core_properties
    return {
        "title": props.title,
        "author": props.author,
        "created": props.created,
    }
